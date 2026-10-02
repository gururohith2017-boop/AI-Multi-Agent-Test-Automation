from pathlib import Path
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent

REQUIREMENT_FILE = (
    PROJECT_ROOT
    / "user_stories"
    / "customer_support_user_stories.xlsx"
)


def load_requirements():

    if not REQUIREMENT_FILE.exists():

        raise FileNotFoundError(
            f"Requirement file not found: {REQUIREMENT_FILE}"
        )

    workbook = load_workbook(
        REQUIREMENT_FILE,
        data_only=True
    )

    worksheet = workbook["Requirements"]

    requirements = []

    headers = [
        cell.value
        for cell in worksheet[1]
    ]

    for row in worksheet.iter_rows(
        min_row=2,
        values_only=True
    ):

        if not row[0]:
            continue

        requirement = dict(
            zip(headers, row)
        )

        requirements.append(
            requirement
        )

    workbook.close()

    return requirements


if __name__ == "__main__":

    requirements = load_requirements()

    print()
    print("==============================================")
    print("REQUIREMENT READER")
    print("==============================================")
    print()

    print(
        f"Requirement File: {REQUIREMENT_FILE}"
    )

    print(
        f"Total Test Cases: {len(requirements)}"
    )

    print()

    for requirement in requirements:

        print(
            f"{requirement['Test Case ID']} "
            f"- "
            f"{requirement['Test Scenario']}"
        )

    print()
    print("Requirement loading completed successfully.")